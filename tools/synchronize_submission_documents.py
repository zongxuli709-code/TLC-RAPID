"""Synchronize the manuscript and supplement with TLC-RAPID v1.0 behavior."""

import subprocess
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
JCA = ROOT / "release" / "JCA-20260908"
RELEASE_URL = "https://github.com/zongxuli709-code/TLC-RAPID/releases/tag/v1.0"
WEIGHTS_URL = "https://github.com/zongxuli709-code/TLC-RAPID/releases/download/v1.0/best.pt"


def current_commit() -> str:
    completed = subprocess.run(
        ["git", "rev-parse", "HEAD"],
        cwd=ROOT,
        check=True,
        capture_output=True,
        text=True,
    )
    return completed.stdout.strip()


def replace_paragraph_starting_with(document: Document, prefix: str, replacement: str) -> None:
    matches = [paragraph for paragraph in document.paragraphs if paragraph.text.strip().startswith(prefix)]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one paragraph beginning {prefix!r}; found {len(matches)}")
    matches[0].text = replacement


def main() -> None:
    commit = current_commit()
    manuscript_path = JCA / "Manuscript_JCA.docx"
    manuscript = Document(manuscript_path)
    replace_paragraph_starting_with(
        manuscript,
        "The applied amount of the reference standard per band",
        "The applied amount of the reference standard per band, denoted as x (μg band⁻¹), was used as the independent variable, and the corresponding image response, denoted as y, was used as the dependent variable. A quadratic polynomial equation (y = ax² + bx + c) was employed to establish the quantitative calibration model, and the coefficient of determination (R²) was used to evaluate goodness of fit. The measured image response of each sample band was substituted into the corresponding quadratic regression equation for back-calculation. When two finite non-negative quadratic roots within the entered working interval were available, TLC-RAPID applied a predefined reference-standard-guided branch-selection procedure. The measured response–amount pairs of the reference standards on the same plate were ordered by response; duplicate response values were represented by the mean of their corresponding applied amounts. A local amount estimate was obtained by linear interpolation of amount against this response-ordered standard series, and the candidate root closest to this estimate was selected automatically. The same rule was applied to all sample bands without reference to HPLC results or manual root selection. If no admissible quadratic root was available, the program applied linear regression to the same plate-specific standards. If linear inversion was unavailable or did not yield a finite non-negative amount, the response-ordered interpolation estimate was exported as an explicitly identified interpolation fallback. The selected applied amount was subsequently converted into marker-compound content according to the sample weight, extraction volume, dilution factor, and actual application volume.",
    )
    replace_paragraph_starting_with(
        manuscript,
        "The measured response of each sample band is substituted into the corresponding quadratic equation for back-calculation.",
        "The measured response of each sample band was substituted into the corresponding quadratic equation for back-calculation. Because a fitted quadratic curve may contain a turning point within the entered standard interval, inverse calculation can yield two non-negative candidate amounts for a given response. In this situation, TLC-RAPID applied a predefined reference-standard-guided branch-selection procedure. The measured response–amount pairs of the reference standards on the same plate were ordered by response. Where duplicate response values occurred, their corresponding applied amounts were averaged. A local amount estimate was then obtained by linear interpolation of amount against the response-ordered standard series. The program automatically selected the candidate quadratic root within the entered working interval that was closest to this local estimate. The same computational rule was applied consistently to all sample bands without using HPLC reference results or operator-selected expected concentrations. If no admissible quadratic root was available, TLC-RAPID applied linear regression to the same plate-specific standard series. If linear inversion was unavailable or did not yield a finite non-negative amount, the response-ordered interpolation estimate was exported as an interpolation fallback. The applied method, root-selection reason, fallback reason, and range status were saved in the exported output.",
    )
    replace_paragraph_starting_with(
        manuscript,
        "The detected bounding boxes, band numbers, quantitative image features",
        "The detected bounding boxes, band numbers, quantitative image features, regression parameters, quadratic candidate roots, local interpolation estimates, selected solutions, back-calculated amounts, applied methods, fallback reasons, range assessments, and summary tables were automatically saved and exported. The results reported in this study were generated using the plate-specific calibration workflow recorded for each analysis. Manual selection of quadratic roots was not performed.",
    )
    replace_paragraph_starting_with(
        manuscript,
        "A quadratic polynomial model was fitted independently to the reference-standard responses on each actual-sample plate.",
        "A quadratic polynomial model was fitted independently to the reference-standard responses on each actual-sample plate. The measured response of each sample band was then substituted into the corresponding plate-specific quadratic equation for back-calculation. When the equation yielded a single finite non-negative candidate solution within the entered working interval, that solution was retained. When two finite non-negative roots within the interval were obtained, TLC-RAPID applied the predefined reference-standard-guided branch-selection procedure described in Section 2.5. The same-plate standard response–amount pairs were ordered by response, and a local amount estimate was obtained by linear interpolation of amount against that ordered standard series; duplicate response values were represented by their mean applied amount. The candidate quadratic root closest to this estimate was selected automatically. This rule was applied consistently to all sample bands without manual root selection, operator-entered expected sample concentrations, or access to the corresponding HPLC results.",
    )
    replace_paragraph_starting_with(
        manuscript,
        "If quadratic inversion did not produce an admissible finite non-negative solution",
        "If quadratic inversion did not produce an admissible finite non-negative solution within the entered working interval, the software applied linear regression to the same plate-specific reference-standard series as a fallback method. If linear inversion was unavailable or did not yield a finite non-negative result, the response-ordered interpolation estimate was exported as an interpolation fallback. The requested calibration model, the method actually applied, the selected result, and the reason for any fallback were recorded in the exported output. Results obtained using fallback methods were explicitly identified and were not represented as direct quadratic back-calculations. Sample responses outside the measured response range of the reference standards, or back-calculated amounts outside the entered working interval, were flagged as out-of-range estimates and were not considered validated quantitative determinations without further verification.",
    )
    replace_paragraph_starting_with(
        manuscript,
        "The source code and trained model weights for TLC-RAPID V1.0",
        f"The source code and trained model weights for TLC-RAPID V1.0 are available in the immutable GitHub V1.0 release ({RELEASE_URL}), which corresponds to commit {commit}. The release archive contains five worked TLC example images and their image-specific reference-standard values; release verification produced numerical results for every detected sample band in all five image sets. The source code is licensed under the GNU Affero General Public License v3.0 (AGPL-3.0). The data supporting the findings of this study are available within the article and its Supplementary Information. Additional data are available from the corresponding author upon reasonable request. Supplementary Software S1 contains the standalone application, while the source code and model weights are distributed through the same GitHub V1.0 release.",
    )
    manuscript_output = JCA / "Manuscript_JCA_v1.0_final.docx"
    manuscript.save(manuscript_output)

    supplement_path = JCA / "Supplementary material.docx"
    supplement = Document(supplement_path)
    replace_paragraph_starting_with(
        supplement,
        "TLC-RAPID (Recognition-based Automated Planar Image Densitometry) V1.0",
        f"TLC-RAPID (Recognition-based Automated Planar Image Densitometry) V1.0 was developed in Python 3.12.4 and integrates a YOLOv5-based target-band segmentation model with OpenCV-based image processing and quantitative feature extraction. The standalone user package supports 64-bit Windows 10 and Windows 11 and does not require a separate Python installation. Supplementary Software S1 is this Windows user package and includes the executable, trained model weight, dataset and class configuration, input templates, dependency information, and user guide. The immutable GitHub V1.0 release ({RELEASE_URL}) corresponds to commit {commit}. The package includes five worked TLC example images named Arnebiae Radix.jpg, Phellodendri Chinensis Cortex.jpg, Lonicerae Flos.jpg, Lonicerae Japonicae Flos.jpg, and Gentianae Macrophyllae Radix.jpg, together with their per-image reference-standard values. Release verification produced a numerical result for every detected sample band across all five example image sets; the exported table identifies direct quadratic results, branch-selected quadratic results, and any fallback estimates. These examples demonstrate software operation and do not replace the formal validation datasets.",
    )
    replace_paragraph_starting_with(
        supplement,
        "The measured response of each sample band was substituted into the corresponding plate-specific quadratic equation",
        "The measured response of each sample band was substituted into the corresponding plate-specific quadratic equation for back-calculation. When quadratic inversion yielded two finite non-negative candidate roots within the entered working interval, TLC-RAPID applied a predefined reference-standard-guided branch-selection procedure. The measured response–amount pairs of the standards on the same plate were ordered by response. Duplicate response values were represented by the mean of their corresponding applied amounts. A local amount estimate was calculated by linear interpolation of amount against this response-ordered standard series, and the candidate quadratic root closest to that estimate was selected automatically. The same computational rule was applied to all sample bands without manual root selection, operator-entered expected sample concentrations, or access to HPLC reference results. If no admissible quadratic root was available, the software used linear regression on the same plate-specific standards; if linear inversion was unavailable or did not yield a finite non-negative amount, the response-ordered interpolation estimate was exported as an explicitly identified interpolation fallback.",
    )
    replace_paragraph_starting_with(
        supplement,
        "This branch-selection procedure provided an objective and reproducible rule",
        "This branch-selection procedure provided an objective and reproducible computational rule for choosing between candidate quadratic roots. When a fitted quadratic curve contained a turning point within the entered standard interval, however, the inverse relationship was not mathematically unique across the entire interval. The response-ordered interpolation estimate was used only as a predefined reference for branch selection or an explicitly labelled fallback; it did not eliminate the underlying uncertainty associated with a non-monotonic calibration curve. This limitation is considered in the interpretation of the actual-sample results in Section 3.6 of the main text.",
    )
    replace_paragraph_starting_with(
        supplement,
        "The quantitative results reported in this study were obtained using the plate-specific quadratic calibration curves",
        "The quantitative results reported in this study were generated using the plate-specific calibration workflow recorded for each analysis. Quadratic candidate roots, response-ordered local interpolation estimates, selected roots, applied methods, fallback reasons, and range flags were retained in the exported output. Manual selection of quadratic roots was not performed. Band-detection performance was evaluated using an independent test subset comprising 82 TLC images, including 31 visible-light images, 37 images acquired at 366 nm, and 14 images acquired at 254 nm. Model outputs were manually verified against the visually identifiable target bands in each image. Correctly detected target bands were classified as true positives (TP), incorrectly detected non-target regions as false positives (FP), and missed target bands as false negatives (FN). Precision, recall, and F1-score were calculated from the resulting TP, FP, and FN counts. Detailed results are provided in Supplementary Table S1.",
    )
    replace_paragraph_starting_with(
        supplement,
        "When inverse calculation of a plate-specific quadratic equation produced two finite non-negative candidate roots",
        "When inverse calculation of a plate-specific quadratic equation produced two finite non-negative candidate roots within the entered working interval, TLC-RAPID applied the predefined reference-standard-guided branch-selection procedure described in Supplementary Section S5. The same-plate standard response–amount pairs were ordered by response, duplicate responses were represented by their mean applied amount, and a local amount estimate was calculated by linear interpolation of amount against the response-ordered standard series. The candidate quadratic root closest to this local estimate was selected automatically. The same rule was applied to every sample band. Root selection did not involve manual judgment, operator-entered expected sample concentrations, or the corresponding HPLC reference result.",
    )
    replace_paragraph_starting_with(
        supplement,
        "Several plate-specific quadratic curves contained a turning point",
        "Several plate-specific quadratic curves contained a turning point within the entered standard interval. The automatic branch-selection procedure provided a consistent and reproducible computational rule in these cases, although the inverse relationship was not mathematically unique across the entire interval. Results obtained from such curves were therefore interpreted together with the limitations discussed in Section 3.6 of the main text. When no admissible quadratic root was available, any linear or interpolation fallback was explicitly labelled in the output. Sample responses outside the response range of the standards, or selected amounts outside the entered standard interval, were identified as outside the corresponding working calibration range and were not regarded as validated quantitative determinations without further verification.",
    )
    replace_paragraph_starting_with(
        supplement,
        "The package includes the standalone executable program, final trained YOLO model-weight file",
        f"The package includes the standalone executable program, final trained YOLO model-weight file (best.pt), dataset and class-configuration file (boCenColor.yaml), input templates, dependency information, five worked example images, and a README.md user guide. Users place TLC images in the designated input directory, enter the reference-standard concentrations in the supplied CSV template, and launch the executable program. TLC-RAPID automatically exports annotated TLC images, fitted calibration curves, and quantitative result tables. The source code corresponding to this package is publicly available in the immutable TLC-RAPID V1.0 GitHub release ({RELEASE_URL}; commit {commit}) under AGPL-3.0.",
    )
    supplement_output = JCA / "Supplementary material_v1.0_final.docx"
    supplement.save(supplement_output)

    print(f"Created {manuscript_output}")
    print(f"Created {supplement_output}")


if __name__ == "__main__":
    main()
